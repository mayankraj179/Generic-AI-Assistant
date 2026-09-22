"""Test-only helpers that build small, REAL documents of each supported
binary format on disk — not mocks, not static checked-in binary blobs.

Generating them in code (rather than committing .pdf/.docx/.xlsx files)
keeps fixtures reviewable as a diff and easy to tweak, and avoids binary
files in version control. DOCX/XLSX use their own libraries' native writers
(python-docx/openpyxl can both write, not just read). PDF has no writer
dependency in this project, so write_minimal_pdf hand-builds a minimal but
fully valid PDF by hand-writing the standard PDF object/xref/trailer
structure — verified to round-trip correctly through pypdf's own reader
before being used as a test fixture.
"""

from __future__ import annotations

from pathlib import Path


def write_minimal_pdf(path: Path, text: str) -> None:
    """Writes a single-page PDF whose content stream draws ``text`` (one
    line per ``\\n``-separated line) as literal text objects — real,
    extractable text, not an image of text.
    """
    lines = text.split("\n")
    content_lines = []
    y = 750
    for line in lines:
        safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content_lines.append(f"BT /F1 12 Tf 50 {y} Td ({safe}) Tj ET")
        y -= 18
    content_bytes = "\n".join(content_lines).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 4 0 R >> >> "
        b"/MediaBox [0 0 612 792] /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content_bytes)).encode() + b" >>\nstream\n"
        + content_bytes + b"\nendstream",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_offset = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF"
    ).encode()

    path.write_bytes(bytes(out))


def write_minimal_docx(path: Path, *, heading: str, paragraphs: list[str]) -> None:
    from docx import Document

    document = Document()
    document.add_heading(heading, level=1)
    for paragraph_text in paragraphs:
        document.add_paragraph(paragraph_text)
    document.save(str(path))


def write_minimal_xlsx(
    path: Path, *, sheet_title: str, header: list[str], rows: list[list[object]]
) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_title
    worksheet.append(header)
    for row in rows:
        worksheet.append(row)
    workbook.save(str(path))
