from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from html import unescape
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from zipfile import BadZipFile

import docx
import openpyxl
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from openpyxl.utils.exceptions import InvalidFileException as XlsxInvalidFileException
from pypdf import PdfReader
from pypdf.errors import PyPdfError


class ParseError(Exception):
    """Raised when a source file can't be parsed into text — a corrupt or
    unreadable PDF/DOCX/XLSX, for example. Wraps the underlying
    library-specific exception (pypdf/python-docx/openpyxl each raise their
    own types) so callers never see a raw vendor exception, the same pattern
    ModelProviderError/EmbeddingProviderError already establish elsewhere in
    this codebase.
    """


@dataclass(frozen=True)
class DiscoveredSource:
    """One raw document location found by a source connector."""

    uri: str
    source_type: str  # e.g. "confluence", "sharepoint", "filesystem"


@dataclass(frozen=True)
class SourceWithAccess:
    """A discovered source plus the access labels captured from it,
    before any parsing has happened. Access must be captured at the
    source — never inferred or defaulted after the fact."""

    source: DiscoveredSource
    access_labels: frozenset[str]


@dataclass(frozen=True)
class ParsedDocument:
    """Raw text extracted from a source, still access-labeled.

    ``content_hash`` is the sha256 hex digest of ``raw_text`` (the parsed
    text, not the source file's raw bytes) — deliberately hashing the
    extracted content rather than the file: two files with identical visible
    text but different bytes (e.g. a DOCX re-saved by a different Office
    version touches internal XML/zip metadata even when nothing a reader
    would notice changed) must hash the same, since what actually matters
    for re-ingestion idempotency is "did the content that gets embedded
    change", not "did the file change byte-for-byte".
    """

    source: SourceWithAccess
    title: str
    raw_text: str
    content_hash: str


@dataclass(frozen=True)
class Chunk:
    """A single chunk ready for embedding. embedded_text carries a
    breadcrumb prefix (title/section context); display_text is what
    gets shown to the user, without the breadcrumb noise.

    ``similarity_score`` is populated only when a Chunk comes back from a
    vector search (PgVectorStore.search) — it's ``1 - cosine_distance``, so
    higher is a closer match: 1.0 is identical, 0.0 is orthogonal/unrelated,
    and negative values point in opposite directions. It's ``None`` for
    chunks that were never scored (e.g. freshly produced by the ingestion
    pipeline, before they're embedded and stored, or from a fake retrieval
    service in tests) — callers must treat ``None`` as "not measured, don't
    filter on it" rather than "worst possible score."
    """

    document_title: str
    chunk_index: int
    display_text: str
    embedded_text: str
    access_labels: frozenset[str]
    similarity_score: float | None = None


_PLAIN_TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".json", ".yaml", ".yml", ".log",
}
_HTML_SUFFIXES = {".html", ".htm"}
_PDF_SUFFIXES = {".pdf"}
_DOCX_SUFFIXES = {".docx"}
_XLSX_SUFFIXES = {".xlsx"}

# What discover() will find and parse() knows how to handle. Kept as one set
# (rather than a dict/dispatch table) since discover() only needs membership,
# not per-suffix behavior — parse() below does its own suffix dispatch.
_SUPPORTED_SUFFIXES = (
    _PLAIN_TEXT_SUFFIXES | _HTML_SUFFIXES | _PDF_SUFFIXES | _DOCX_SUFFIXES | _XLSX_SUFFIXES
)

_LABEL_RE = re.compile(r"^[A-Za-z0-9_-]+:[A-Za-z0-9_.-]+$")


def _strip_uri_metadata(uri: str) -> str:
    """Return the underlying filesystem path for a local file URI while dropping
    any query-string metadata used to carry access labels.
    """
    sanitized = uri.split("?", 1)[0].split("#", 1)[0]
    if sanitized.startswith("file://"):
        parsed = urlparse(sanitized)
        return parsed.path or sanitized
    return sanitized


def _iter_label_candidates_from_uri(uri: str) -> set[str]:
    """Capture labels explicitly declared in the source URI path or query string."""
    labels: set[str] = set()

    def _valid_label(value: str) -> bool:
        if not value or value in {".", "..", ":"}:
            return False
        if len(value) >= 2 and value[1] == ":" and value[0].isalpha():
            return False
        return bool(_LABEL_RE.fullmatch(value.strip()))

    try:
        parsed = urlparse(uri)
    except ValueError:
        parsed = None

    if parsed is not None:
        # Typical URI sources may encode access labels as query params, e.g.
        # ?labels=role:employee&labels=dept:hr.
        for values in parse_qs(parsed.query, keep_blank_values=True).values():
            for value in values:
                labels.update(
                    part.strip()
                    for part in value.split(",")
                    if _valid_label(part.strip())
                )
        if parsed.path:
            for segment in Path(parsed.path).parts:
                value = segment.strip().strip("/")
                if _valid_label(value):
                    labels.add(value)

    path = Path(_strip_uri_metadata(uri))
    for segment in path.parts:
        value = segment.strip().strip("/")
        if _valid_label(value):
            labels.add(value)

    return labels


def discover(source_root: str) -> list[DiscoveredSource]:
    """Stage 1: enumerate candidate documents from a source connector."""
    root = Path(source_root)
    if not root.exists():
        raise FileNotFoundError(f"source_root does not exist: {source_root}")

    if root.is_file():
        candidates = [root]
    else:
        candidates = [
            path
            for path in root.rglob("*")
            if path.is_file() and path.suffix.lower() in _SUPPORTED_SUFFIXES
        ]

    return [
        DiscoveredSource(uri=str(path), source_type="filesystem")
        for path in sorted(candidates, key=lambda item: str(item).lower())
    ]


def capture_access(source: DiscoveredSource) -> SourceWithAccess:
    """Stage 2: read the source system's own access control and translate
    it into our namespaced label format. Must happen before parsing —
    never derive access from document content.
    """
    labels = _iter_label_candidates_from_uri(source.uri)
    return SourceWithAccess(source=source, access_labels=frozenset(labels))


def _parse_plain_text(file_path: Path) -> str:
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return file_path.read_text(encoding="utf-8", errors="replace")


def _parse_html(file_path: Path) -> str:
    raw_text = _parse_plain_text(file_path)
    raw_text = re.sub(r"<[^>]+>", " ", raw_text)
    return unescape(raw_text)


def _parse_pdf(file_path: Path) -> str:
    """Extracts text page by page, joined with a paragraph break so
    chunk()'s existing paragraph/sentence splitting still respects natural
    page boundaries. Page numbers are available here (``enumerate(pages)``)
    but deliberately not threaded any further — see the module-level note
    in IngestionService/this module's docstring in the PR summary: doing so
    properly would mean chunk() tracking per-chunk source offsets, which is
    a chunking-algorithm change out of scope for this pass.
    """
    try:
        reader = PdfReader(str(file_path))
        pages_text = [page.extract_text() or "" for page in reader.pages]
    except PyPdfError as exc:
        raise ParseError(f"could not parse PDF file: {file_path}") from exc
    except Exception as exc:  # some corrupt PDFs fail outside pypdf's own error hierarchy
        raise ParseError(f"could not parse PDF file: {file_path}") from exc
    return "\n\n".join(pages_text)


def _parse_docx(file_path: Path) -> str:
    """Extracts paragraph text, including headings (as plain text — heading
    *styles* aren't carried forward into a breadcrumb path; see this
    module's PR summary for why that was deliberately skipped for this
    pass). Paragraphs are joined with a blank line so heading/section
    transitions become natural chunk()-recognized paragraph breaks even
    without explicit heading metadata.
    """
    try:
        document = docx.Document(str(file_path))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    except DocxPackageNotFoundError as exc:
        raise ParseError(f"could not parse DOCX file: {file_path}") from exc
    except Exception as exc:
        raise ParseError(f"could not parse DOCX file: {file_path}") from exc
    return "\n\n".join(paragraphs)


def _parse_xlsx(file_path: Path) -> str:
    """Extracts each sheet as structured "Row N: col=value, ..." text rather
    than one blob — a leave-balance-style table reads far better this way
    than as a flat wall of cell values, and it's still plain text chunk()
    can split normally.
    """
    try:
        workbook = openpyxl.load_workbook(str(file_path), data_only=True, read_only=True)
    except (BadZipFile, XlsxInvalidFileException) as exc:
        raise ParseError(f"could not parse XLSX file: {file_path}") from exc
    except Exception as exc:
        raise ParseError(f"could not parse XLSX file: {file_path}") from exc

    sheet_blocks: list[str] = []
    for worksheet in workbook.worksheets:
        header: tuple | None = None
        row_lines: list[str] = []
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if row_index == 1:
                header = row
                continue
            if all(cell is None for cell in row):
                continue
            if header:
                cells = ", ".join(
                    f"{col}={value}"
                    for col, value in zip(header, row, strict=False)
                    if col is not None and value is not None
                )
            else:
                cells = ", ".join(str(value) for value in row if value is not None)
            if cells:
                row_lines.append(f"Row {row_index}: {cells}")
        if row_lines:
            sheet_blocks.append(f"Sheet: {worksheet.title}\n" + "\n".join(row_lines))

    return "\n\n".join(sheet_blocks)


def parse(source: SourceWithAccess) -> ParsedDocument:
    """Stage 3: extract text from the source's native format.

    Dispatches by file extension to the plain-text path (unchanged
    behavior), the HTML path (unchanged, still strips tags), or one of the
    new PDF/DOCX/XLSX paths. Any format-specific parsing failure is wrapped
    into ParseError — never a raw pypdf/python-docx/openpyxl exception.
    """
    file_path = Path(_strip_uri_metadata(source.source.uri))
    if not file_path.exists() or not file_path.is_file():
        raise FileNotFoundError(f"source file does not exist: {file_path}")

    return parsed_document(
        source, title=file_path.stem or file_path.name, raw_text=extract_text(file_path)
    )


def is_supported_file(name: str) -> bool:
    """Whether discover()/extract_text() handle this file name's suffix."""
    return Path(name).suffix.lower() in _SUPPORTED_SUFFIXES


def extract_text(file_path: Path) -> str:
    """parse()'s format dispatch on its own, for knowledge-source adapters
    whose files aren't addressed by a local path URI (e.g. an object
    downloaded to a temp file)."""
    suffix = file_path.suffix.lower()
    if suffix in _PDF_SUFFIXES:
        return _parse_pdf(file_path)
    if suffix in _DOCX_SUFFIXES:
        return _parse_docx(file_path)
    if suffix in _XLSX_SUFFIXES:
        return _parse_xlsx(file_path)
    if suffix in _HTML_SUFFIXES:
        return _parse_html(file_path)
    return _parse_plain_text(file_path)


def parsed_document(source: SourceWithAccess, *, title: str, raw_text: str) -> ParsedDocument:
    """Builds a ParsedDocument with the one content-hash rule (sha256 of the
    stripped text) every source shares, so idempotency means the same thing
    for a file, an object and a database row."""
    stripped_text = raw_text.strip()
    content_hash = hashlib.sha256(stripped_text.encode("utf-8")).hexdigest()
    return ParsedDocument(
        source=source, title=title, raw_text=stripped_text, content_hash=content_hash
    )


def chunk(document: ParsedDocument) -> list[Chunk]:
    """Stage 4: split parsed text into embedding-sized chunks, carrying
    the document's access labels onto every chunk."""
    raw_text = document.raw_text.strip()
    if not raw_text:
        return []

    max_chars = 800
    segments = re.split(r"\n\s*\n+|\r\n\s*\r\n+", raw_text)
    chunks: list[str] = []

    for segment in segments:
        cleaned = segment.strip()
        if not cleaned:
            continue
        sentences = re.split(r"(?<=[.!?])\s+", cleaned)
        current = ""
        for sentence in sentences:
            candidate = sentence if not current else f"{current} {sentence}"
            if len(candidate) <= max_chars or not current:
                current = candidate
            else:
                chunks.append(current.strip())
                current = sentence
        if current.strip():
            chunks.append(current.strip())

    if not chunks:
        chunks = [raw_text]

    final_chunks: list[Chunk] = []
    for index, text in enumerate(chunks):
        display_text = text.strip()
        embedded_text = f"{document.title}: {display_text}"
        final_chunks.append(
            Chunk(
                document_title=document.title,
                chunk_index=index,
                display_text=display_text,
                embedded_text=embedded_text,
                access_labels=document.source.access_labels,
            )
        )
    return final_chunks


def run_pipeline(source_root: str) -> list[Chunk]:
    """Wires the four stages together end-to-end.

    Embedding + atomic commit (two-phase, generation-id based visibility)
    is a later stage — deliberately not in scope for Day 1.
    """
    all_chunks: list[Chunk] = []
    for discovered in discover(source_root):
        with_access = capture_access(discovered)
        parsed = parse(with_access)
        all_chunks.extend(chunk(parsed))
    return all_chunks
